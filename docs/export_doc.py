import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def markdown_to_docx(md_file_path, docx_file_path):
    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Set standard margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_code_block():
        nonlocal code_lines
        if code_lines:
            code_text = "".join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(code_text.rstrip())
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(40, 40, 40)
            # Add subtle gray background box
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F7"/>')
            pPr.append(shd)
            code_lines = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            col_count = max(len(r) for r in table_rows)
            table = doc.add_table(rows=len(table_rows), cols=col_count)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            for row_idx, row_data in enumerate(table_rows):
                is_header = (row_idx == 0)
                row = table.rows[row_idx]
                for col_idx in range(col_count):
                    val = row_data[col_idx] if col_idx < len(row_data) else ""
                    cell = row.cells[col_idx]
                    cell.text = val
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        for r in p.runs:
                            r.font.name = 'Arial'
                            r.font.size = Pt(9.5)
                            if is_header:
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(30, 41, 59)
                    if is_header:
                        set_cell_background(cell, "E2E8F0")
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")

            doc.add_paragraph() # Spacing
            table_rows = []

    for line in lines:
        stripped = line.strip()

        # Handle Code Block fences
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code_block()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Handle Table
        if stripped.startswith("|") and stripped.endswith("|"):
            # Check separator line
            if re.match(r'^\|[\s:-]+\|$', stripped) or ":---" in stripped or "---:" in stripped:
                continue
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(cells)
            continue
        else:
            if in_table:
                in_table = False
                flush_table()

        # Empty line
        if not stripped:
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[3:])
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 116, 144)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            run = h.add_run(stripped[4:])
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 41, 55)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            run = h.add_run(stripped[5:])
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(55, 65, 81)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            # Process bold runs
            raw_text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', raw_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            raw_text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', raw_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif stripped.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            pPr.append(pBdr)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    doc.save(docx_file_path)
    print(f"Successfully generated {docx_file_path}")

if __name__ == "__main__":
    md_path = r"E:\e-commerce-microservices-master\e-commerce-microservices-master\DANH_SACH_CHI_TIET_API.md"
    docx_path = r"E:\e-commerce-microservices-master\e-commerce-microservices-master\TAI_LIEU_CHI_TIET_API_VA_CONG_NGHE.docx"
    markdown_to_docx(md_path, docx_path)
